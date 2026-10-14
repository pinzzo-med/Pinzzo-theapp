"""Merge laptop + cloud per-store stock with the store map and catalogue into
Plazza_Store_Inventory.xlsx and plazza_store_stock_long.csv.zip."""
import csv, json, sqlite3, statistics, sys, zipfile
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
UPLOAD = Path(sys.argv[1])  # folder with the laptop store_stock.csv / store_summary.csv
OUT = HERE / 'out'
OUT.mkdir(exist_ok=True)
csv.field_size_limit(10**8)

LABEL = {'4436157a': 'Brookefield / Mahadevapura', '7e4b6aef': 'Chandapura / Bommasandra',
         '7e91be9f': 'Kalkere / Begur / Hulimavu', 'a258ba87': 'Varthur / Gunjur / Panathur',
         'aa49e0e4': 'Kalyan Nagar / Hennur', 'b666cfc6': 'Bellandur / Kodathi',
         'bb28573b': 'HSR / BTM / Koramangala', 'f7070977': 'Basavanagudi / Vijayanagar',
         'fe4c5211': 'Anjanapura / JP Nagar'}

# ---- catalogue (product master)
cat = {}
for pid, raw in sqlite3.connect(HERE / 'plazza.sqlite').execute('SELECT product_id, raw FROM products'):
    cat[pid] = json.loads(raw)

# ---- stock: store -> product -> record
stock = defaultdict(dict)
searches = {}
cloud = sqlite3.connect(HERE / 'stock.sqlite')
for sid in ('4436157a-bf85-4132-9f8d-c2b5afc4296c', '7e4b6aef-936f-4d8d-93ed-b74892c388aa'):
    searches[sid] = cloud.execute('SELECT count(*) FROM done WHERE store_id=?', (sid,)).fetchone()[0]
    for pid, inv, dt, sp, mrp, disc, offer, raw, at in cloud.execute(
            'SELECT product_id, store_inventory, delivery_type, selling_price, mrp, discount, offer_price, raw, at '
            'FROM stock WHERE store_id=?', (sid,)):
        cat.setdefault(pid, json.loads(raw))
        stock[sid][pid] = dict(inv=inv, dt=dt, sp=sp, mrp=mrp, offer=offer, at=at, source='cloud')

num = lambda v: float(v) if v not in (None, '') else None
for r in csv.DictReader((UPLOAD / 'store_stock.csv').open(encoding='utf-8-sig')):
    sid, pid = r['store_id'], r['product_id']
    stock[sid][pid] = dict(inv=num(r['store_inventory']), dt=r['delivery_type'], sp=num(r['selling_price']),
                           mrp=num(r['mrp']), offer=num(r['offer_price_999']), at=r['collected_at'], source='laptop')
    cat.setdefault(pid, {'product_id': pid, 'name': r['name'], 'brand': r['brand'], 'manufacturers': r['manufacturers'],
                         'salt_composition': r['salt_composition'], 'packaging_detail': r['packaging_detail'],
                         'tag': r['tag'], 'prescription_required': r['prescription_required'] == 'True'})
for r in csv.DictReader((UPLOAD / 'store_summary.csv').open(encoding='utf-8-sig')):
    if r['store_id'] not in searches:
        searches[r['store_id']] = int(r['searches_done'].split('/')[0])
TOTAL_Q = 2659

# ---- store map
loc = sqlite3.connect(HERE / 'stores.sqlite')
pts = defaultdict(list)
for row in loc.execute('SELECT store_id, pincode, locality, eta_min, distance_m FROM locations WHERE serviceable=1'):
    pts[row[0]].append(row[1:])

stores = sorted(stock, key=lambda s: LABEL.get(s[:8], s))
lab = lambda s: f'{LABEL.get(s[:8], s[:8])} ({s[:8]})'


def group(tag):
    t = (tag or '').lower()
    if t == 'prescription required': return 'Rx medicines'
    if t == 'drugs': return 'Medicines (no Rx flag)'
    if t == 'otc': return 'OTC'
    if any(k in t for k in ('ayurved', 'herbal', 'unani', 'homeo', 'siddha')): return 'Ayurveda / herbal'
    if any(k in t for k in ('supplement', 'nutrition', 'vitamin', 'protein')): return 'Supplements'
    if t in ('general', ''): return 'General / personal care'
    return 'Other'


wb = Workbook()
HF, HFILL = Font(bold=True, color='FFFFFF'), PatternFill('solid', start_color='0A3A75')


def sheet(title, headers, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HF, HFILL
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(11, min(len(str(h)) + 3, 28)))
    ws.freeze_panes = 'B2'
    return ws


allpids = set(cat) | {p for s in stock for p in stock[s]}
in_stock_any = {p for s in stock for p, v in stock[s].items() if (v['inv'] or 0) > 0}
dates = sorted(v['at'] for s in stock for v in stock[s].values() if v['at'])
sheet('Summary', ['Item', 'Value'], [
    ('What this is', 'Plazza (Bengaluru) product, price and stock per dark store, from the catalogue API their website uses'),
    ('Collected (UTC)', f'{dates[0]} to {dates[-1]}'),
    ('Stores', len(stores)),
    ('Pincodes served (from 729-point grid)', len({p[0] for s in pts for p in pts[s] if p[0]})),
    ('Products in catalogue', len(allpids)),
    ('Products in stock in at least one store', len(in_stock_any)),
    ('Store x product rows', sum(len(v) for v in stock.values())),
    ('Stock meaning', 'store_inventory = quantity Plazza\'s API reported for that store at collection time. A snapshot of stock, not sales.'),
    ('Completeness', 'See Stores sheet: Varthur (a258ba87) ran 472 of 2,659 searches; Brookefield and Chandapura ran the earlier 1,929-search list. Missing products for those stores are "not seen", not "out of stock".'),
    ('Blank in matrix', 'Product not seen at that store (search not run or not returned). 0 = listed but none in store.'),
], {'Item': 38, 'Value': 110}, first=True)

srows = []
for s in stores:
    p = pts.get(s, [])
    etas = [x[2] for x in p if x[2] is not None]
    recs = stock[s].values()
    ins = sum(1 for v in recs if (v['inv'] or 0) > 0)
    srows.append([lab(s), s, len({x[0] for x in p if x[0]}), ', '.join(sorted({x[0] for x in p if x[0]})),
                  ', '.join(sorted({x[1] for x in p if x[1]}))[:300], len(p),
                  min(etas) if etas else None, statistics.median(etas) if etas else None, max(etas) if etas else None,
                  round(statistics.mean([x[3] for x in p if x[3]]) / 1000, 1) if p else None,
                  f'{searches.get(s, 0)}/{TOTAL_Q}', len(stock[s]), ins, round(ins / len(stock[s]) * 100, 1),
                  int(sum(v['inv'] for v in recs if (v['inv'] or 0) > 0))])
sheet('Stores', ['Store', 'Store ID', 'Pincodes', 'Pincode list', 'Localities', 'Grid points served', 'ETA min (min)',
                 'ETA median', 'ETA max', 'Avg distance km', 'Searches run', 'Products seen', 'In stock',
                 'In-stock %', 'Units in stock'], srows, {'Store': 34, 'Store ID': 12, 'Pincode list': 40, 'Localities': 50})

pin = defaultdict(lambda: defaultdict(list))
for s in pts:
    for pc, locality, eta, dist in pts[s]:
        if pc:
            pin[pc][s].append((eta, dist, locality))
prow = []
for pc in sorted(pin):
    for s, v in sorted(pin[pc].items(), key=lambda kv: -len(kv[1])):
        etas = [x[0] for x in v if x[0] is not None]
        prow.append([pc, lab(s), len(v), min(etas) if etas else None, max(etas) if etas else None,
                     round(min(x[1] for x in v if x[1]) / 1000, 1), ', '.join(sorted({x[2] for x in v if x[2]}))[:200]])
sheet('Pincodes', ['Pincode', 'Served by store', 'Grid points', 'ETA min', 'ETA max', 'Nearest km', 'Localities'], prow,
      {'Served by store': 36, 'Localities': 60})

# availability by category x store
acat = defaultdict(lambda: defaultdict(lambda: [0, 0]))
for s in stores:
    for pid, v in stock[s].items():
        g = group(cat.get(pid, {}).get('tag'))
        acat[g][s][0] += 1
        acat[g][s][1] += (v['inv'] or 0) > 0
sheet('In-stock % by category', ['Category'] + [lab(s) for s in stores],
      [[g] + [round(acat[g][s][1] / acat[g][s][0] * 100, 1) if acat[g][s][0] else None for s in stores]
       for g in sorted(acat, key=lambda g: -sum(x[0] for x in acat[g].values()))],
      {'Category': 26})

# stock matrix
heads = ['Product ID', 'Name', 'Brand', 'Manufacturer', 'Salt', 'Pack', 'Category', 'Rx', 'MRP', 'Price (min)',
         'Price (max)', 'Stores in stock', 'Total units'] + [lab(s) for s in stores]
mrows = []
for pid in allpids:
    c = cat.get(pid, {})
    recs = [stock[s].get(pid) for s in stores]
    prices = [r['sp'] for r in recs if r and r['sp'] is not None]
    mrp = next((r['mrp'] for r in recs if r and r['mrp']), c.get('mrp'))
    mrows.append([pid, c.get('name'), c.get('brand'), c.get('manufacturers'), c.get('salt_composition'),
                  c.get('packaging_detail'), group(c.get('tag')), bool(c.get('prescription_required')), mrp,
                  min(prices) if prices else c.get('plazza_selling_price_incl_gst'),
                  max(prices) if prices else c.get('plazza_selling_price_incl_gst'),
                  sum(1 for r in recs if r and (r['inv'] or 0) > 0),
                  int(sum(r['inv'] for r in recs if r and (r['inv'] or 0) > 0))] +
                 [r['inv'] if r else None for r in recs])
mrows.sort(key=lambda r: (-r[12], r[1] or ''))
sheet('Stock matrix', heads, mrows, {'Name': 42, 'Salt': 32, 'Manufacturer': 24, 'Pack': 22})

top = [r for r in mrows if r[11] >= 1][:1000]
sheet('Deepest stock (top 1000)', heads[:13], [r[:13] for r in top], {'Name': 42, 'Salt': 32, 'Manufacturer': 24, 'Pack': 22})

pv = [r[:11] + [round((r[10] - r[9]) / r[10] * 100, 1)] for r in mrows
      if r[9] is not None and r[10] is not None and r[10] > r[9] and r[10] > 0]
pv.sort(key=lambda r: -r[11])
sheet('Price differs by store', heads[:11] + ['Spread %'], pv, {'Name': 42, 'Salt': 32, 'Manufacturer': 24, 'Pack': 22})

full = [s for s in stores if searches.get(s, 0) >= 1900]
gaps = [r[:13] for r in mrows if all(stock[s].get(r[0]) is not None for s in full) and r[11] == 0]
sheet('Listed, stocked nowhere', heads[:13], gaps, {'Name': 42, 'Salt': 32, 'Manufacturer': 24, 'Pack': 22})

wb.save(OUT / 'Plazza_Store_Inventory.xlsx')

with zipfile.ZipFile(OUT / 'plazza_store_stock_long.csv.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    with z.open('plazza_store_stock_long.csv', 'w') as f:
        import io
        t = io.TextIOWrapper(f, encoding='utf-8-sig', newline='')
        w = csv.writer(t)
        w.writerow(['store_id', 'store', 'product_id', 'name', 'category', 'mrp', 'selling_price', 'offer_price_999',
                    'store_inventory', 'delivery_type', 'collected_at', 'source'])
        for s in stores:
            for pid, v in stock[s].items():
                c = cat.get(pid, {})
                w.writerow([s, LABEL.get(s[:8]), pid, c.get('name'), group(c.get('tag')), v['mrp'], v['sp'], v['offer'],
                            v['inv'], v['dt'], v['at'], v['source']])
        t.flush()
        t.detach()

print('stores', len(stores), 'products', len(allpids), 'in stock somewhere', len(in_stock_any),
      'price-varying', len(pv), 'stocked nowhere', len(gaps))
for r in srows:
    print(r[0], '| pincodes', r[2], '| ETA', r[6], r[7], r[8], '| km', r[9], '|', r[10], '| seen', r[11], 'in', r[12], f'({r[13]}%)', 'units', r[14])
