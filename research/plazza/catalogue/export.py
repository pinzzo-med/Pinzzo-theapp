"""Export plazza.sqlite to Plazza_Catalogue.csv and Plazza_Catalogue.xlsx."""
import csv, json, sqlite3, statistics
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
db = sqlite3.connect(HERE / 'plazza.sqlite')

COLS = ['product_id', 'category_group', 'name', 'brand', 'manufacturers', 'salt_composition', 'product_form', 'package',
        'packaging_detail', 'product_qty', 'product_unit', 'medicine_type', 'tag', 'prescription_required',
        'mrp', 'plazza_selling_price_incl_gst', 'discount_pct', 'effective_customer_discount',
        'offer_price_999', 'offer_text', 'availability', 'store_inventory', 'inventory_quantity',
        'delivery_type', 'fast_delivery', 'deferred_allowed', 'is_active', 'is_banned', 'is_store_only',
        'telemedicine_allowed', 'most_bought', 'image_url', 'listed_as', 'first_seen', 'last_seen']


def availability(p, listed_as):
    if not p.get('is_active', True) or p.get('is_banned'):
        return 'inactive'
    qty = p.get('store_inventory', p.get('inventory_quantity'))
    if isinstance(qty, (int, float)) and qty > 0:
        return 'in_stock'
    # 0 in the store: Plazza either takes the order for later delivery or shows it as out of stock
    return 'deferred_delivery' if p.get('deferred_allowed') else 'out_of_stock'


def group(tag):
    t = (tag or '').lower()
    if t == 'prescription required':
        return 'Rx medicines'
    if t == 'drugs':
        return 'Medicines (no Rx flag)'
    if t == 'otc':
        return 'OTC'
    if any(k in t for k in ('ayurved', 'herbal', 'unani', 'homeo', 'siddha')):
        return 'Ayurveda / herbal'
    if any(k in t for k in ('supplement', 'nutrition', 'vitamin', 'protein')):
        return 'Supplements'
    if t in ('general', ''):
        return 'General / personal care'
    return 'Other'


rows = []
for pid, raw, listed_as, first, last in db.execute('SELECT product_id, raw, listed_as, first_seen, last_seen FROM products'):
    p = json.loads(raw)
    mrp, sp = p.get('mrp'), p.get('plazza_selling_price_incl_gst')
    di = p.get('discount_info') or {}
    r = {k: p.get(k) for k in COLS}
    r.update(
        discount_pct=round((mrp - sp) / mrp * 100, 2) if mrp and sp is not None and mrp > 0 else None,
        offer_price_999=di.get('price'), offer_text=di.get('description'),
        availability=availability(p, listed_as), category_group=group(p.get('tag')),
        image_url=(p.get('image_url') or [None])[0], listed_as=listed_as, first_seen=first, last_seen=last)
    rows.append(r)
rows.sort(key=lambda r: (r['category_group'], r['name'] or ''))

with (HERE / 'Plazza_Catalogue.csv').open('w', newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=COLS)
    w.writeheader()
    w.writerows(rows)

# ---- summaries
n = len(rows)
avail = Counter(r['availability'] for r in rows)
by_tag = defaultdict(list)
for r in rows:
    by_tag[r['category_group']].append(r)


def stats(rs):
    d = [r['discount_pct'] for r in rs if r['discount_pct'] is not None]
    a = Counter(r['availability'] for r in rs)
    known = a['in_stock'] + a['out_of_stock'] + a['deferred_delivery']
    return [len(rs), a['in_stock'], a['deferred_delivery'], a['out_of_stock'], a['inactive'],
            round(a['in_stock'] / known * 100, 1) if known else None,
            round(statistics.median(d), 1) if d else None, round(statistics.mean(d), 1) if d else None,
            sum(1 for r in rs if r['prescription_required']), sum(1 for r in rs if r['offer_price_999'])]


queries = db.execute('SELECT count(*), sum(pass=1), sum(pass=2), sum(status!="ok") FROM queries').fetchone()
errors = db.execute('SELECT count(*) FROM errors').fetchone()[0]
first_seen = db.execute('SELECT min(first_seen), max(last_seen) FROM products').fetchone()

wb = Workbook()
HF, HFILL = Font(bold=True, color='FFFFFF'), PatternFill('solid', start_color='0A3A75')


def table(ws, headers, data, widths=None):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font, c.fill = HF, HFILL
    for row in data:
        ws.append(row)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(len(str(h)) + 4, 30)))
    ws.freeze_panes = ws.cell(2, 1)


ws = wb.active
ws.title = 'Summary'
summary = [
    ('Source', 'cma.plazza.in catalogue search API (what the plazza.in website calls)'),
    ('Location context', 'One guest session with no address set: Plazza default store (Bengaluru). Stock is for that store'),
    ('Collected (UTC)', f'{first_seen[0]} to {first_seen[1]}'),
    ('Unique products', n),
    ('In stock now (store_inventory > 0)', avail['in_stock']),
    ('Deferred delivery (0 in store, can be ordered for later)', avail['deferred_delivery']),
    ('Out of stock', avail['out_of_stock']),
    ('Inactive / banned', avail['inactive']),
    ('Prescription required', sum(1 for r in rows if r['prescription_required'])),
    ('Median discount % vs MRP', round(statistics.median([r['discount_pct'] for r in rows if r['discount_pct'] is not None]), 1)),
    ('Products with "₹999+ order" extra offer', sum(1 for r in rows if r['offer_price_999'])),
    ('Search queries run', f'{queries[0]} (prefix pass {queries[1]}, snowball {queries[2]}, failed {queries[3]})'),
    ('Request errors logged', errors),
    ('Coverage note', 'Search returns ~100 results per query, so coverage comes from many queries. Plazza claims 20,000–100,000+ products; the gap is products no query reached.'),
    ('Stock note', 'store_inventory is the quantity Plazza\'s own API returns for the default store at collection time. It is a snapshot, not sales.'),
]
table(ws, ['Item', 'Value'], summary, {'Item': 48, 'Value': 100})

table(wb.create_sheet('By category'),
      ['Category group', 'Products', 'In stock', 'Deferred', 'Out of stock', 'Inactive', 'In-stock % (of known)',
       'Median discount %', 'Avg discount %', 'Rx required', 'With ₹999+ offer'],
      [[t] + stats(rs) for t, rs in sorted(by_tag.items(), key=lambda kv: -len(kv[1]))])

mf = Counter((r['manufacturers'] or '(none)').strip().upper() for r in rows)
table(wb.create_sheet('Manufacturers'), ['Manufacturer', 'Products'], mf.most_common(300), {'Manufacturer': 50})

bands = [(0, 0), (0.01, 5), (5.01, 10), (10.01, 15), (15.01, 20), (20.01, 30), (30.01, 50), (50.01, 100)]
table(wb.create_sheet('Discount bands'), ['Discount % band', 'Products'],
      [(f'{a}–{b}%' if b else '0%', sum(1 for r in rows if r['discount_pct'] is not None and a <= r['discount_pct'] <= b)) for a, b in bands])

table(wb.create_sheet('Catalogue'), COLS, [[r[c] for c in COLS] for r in rows], {'name': 45, 'salt_composition': 40, 'manufacturers': 28, 'packaging_detail': 28})

wb.save(HERE / 'Plazza_Catalogue.xlsx')
print(f'{n} products exported; availability {dict(avail)}')
