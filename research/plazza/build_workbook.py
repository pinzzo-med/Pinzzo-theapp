"""Build Plazza_Competitive_Intelligence.xlsx: Phase 1 facts pre-filled, the rest as a
structured manual-observation workbook with formulas and dropdowns."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / 'Plazza_Competitive_Intelligence.xlsx'
OBS = '2026-10-12'
HDR_FONT = Font(bold=True, color='FFFFFF')
HDR_FILL = PatternFill('solid', start_color='0A3A75')
NOTE_FONT = Font(italic=True, color='555555')
FACT_FILL = PatternFill('solid', start_color='EEF8D0')

wb = Workbook()
wb.remove(wb.active)


def sheet(title, note, headers, rows=(), widths=None, blank_rows=0):
    ws = wb.create_sheet(title)
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=13, color='0A3A75')
    ws['A2'] = note
    ws['A2'].font = NOTE_FONT
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(h, max(14, min(len(h) + 4, 40)))
    for r in rows:
        ws.append(list(r))
        for cell in ws[ws.max_row]:
            if cell.value not in (None, ''):
                cell.fill = FACT_FILL
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A5'
    return ws


def dropdown(ws, col, options, first=5, last=2000):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col}{first}:{col}{last}')


STATUS = ['observed', 'claimed', 'not_found', 'untestable', 'pending']
SRC = 'https://plazza.in (schema.org JSON-LD in page HTML)'

# 1. Executive Summary
ws = sheet('Executive Summary',
           'Green rows = facts already captured on ' + OBS + '. Formulas count the manual observation tabs.',
           ['Item', 'Value', 'Kind', 'Source', 'Observed'],
           [
               ('Target', 'Plazza (plazza.in), PLAZZA HEALTH PRIVATE LIMITED', 'observed', SRC, OBS),
               ('Registered/office address', '5th Floor, Tower A, No. 51, Innov8 Mantri Commercio, Bellandur, Bengaluru 560103', 'claimed', SRC, OBS),
               ('Service area', 'Bengaluru only (structured data and FAQ)', 'claimed', SRC, OBS),
               ('Delivery promise', '15 minutes; handling 5–15 min stated in Offer data', 'claimed', SRC, OBS),
               ('Hours', '24/7 (00:00–23:59 every day)', 'claimed', SRC, OBS),
               ('Assortment claim', 'Inconsistent: 20,000+ / 35,000+ / 100,000+ in different blocks', 'claimed', SRC, OBS),
               ('Minimum order value', 'None ("no minimum order value", FAQ)', 'claimed', SRC, OBS),
               ('Shipping rate in Offer data', '₹0', 'claimed', SRC, OBS),
               ('Payment methods', 'COD, cards, UPI, net banking, wallets (FAQ)', 'claimed', SRC, OBS),
               ('App rating (from their page)', '4.1 from 31 ratings; not verified against the stores', 'claimed', SRC, OBS),
               ('Automated catalogue collection', 'Not performed: product data loads only from /api/, which robots.txt disallows', 'observed', 'https://plazza.in/robots.txt', OBS),
               ('SKUs observed (manual)', '=COUNTA(\'Master SKU Catalogue\'!D5:D5000)', 'calculated', 'Master SKU Catalogue', ''),
               ('SKUs with selling price (manual)', '=COUNT(\'Master SKU Catalogue\'!P5:P5000)', 'calculated', 'Master SKU Catalogue', ''),
               ('Observed availability % (of SKUs with status)', '=IFERROR(ROUND(100*COUNTIF(\'Master SKU Catalogue\'!R5:R5000,"in_stock")/(COUNTIF(\'Master SKU Catalogue\'!R5:R5000,"in_stock")+COUNTIF(\'Master SKU Catalogue\'!R5:R5000,"out_of_stock")),1),"n/a")', 'calculated', 'Master SKU Catalogue', ''),
               ('Median discount % (manual)', '=IFERROR(MEDIAN(\'Master SKU Catalogue\'!Q5:Q5000),"n/a")', 'calculated', 'Master SKU Catalogue', ''),
               ('Basket scenarios recorded', '=COUNT(\'Basket Experiments\'!L5:L100)', 'calculated', 'Basket Experiments', ''),
               ('ETA rows (incl. claims)', '=COUNTA(\'Delivery ETA\'!E5:E2000)', 'calculated', 'Delivery ETA', ''),
               ('Search queries tested', '=COUNT(\'Search Tests\'!C5:C200)', 'calculated', 'Search Tests', ''),
           ], {'Item': 38, 'Value': 70, 'Source': 40})

# 2. Category Tree
sheet('Category Tree', 'One row per category/subcategory seen in the app or website navigation.',
      ['Category ID (if shown)', 'Category', 'Parent', 'Display order', 'URL / screen', 'Products seen',
       'Available', 'Unavailable', 'Category offer', 'Delivery promise shown', 'Badges', 'Observed at', 'Screenshot', 'Notes'],
      [('', 'Prescription Medicines; OTC; Vitamins & Supplements; Baby Care; Personal Care', '', '', 'homepage structured data', '', '', '', '', '', '', OBS, '', 'Claimed in schema.org ItemList, not navigation'),
       ('', 'Fever; Cold & Cough; Diabetes; Blood Pressure; Antibiotics; Vitamins', '', '', 'homepage structured data', '', '', '', '', '', '', OBS, '', 'Claimed "popular categories" list')])

# 3. Master SKU Catalogue
ws = sheet('Master SKU Catalogue',
           'Record products as an ordinary customer sees them at the fixed address. Keep each pack size as its own row. Leave unknowns blank; never guess.',
           ['Obs ID', 'Observed at (IST)', 'Address PIN', 'Product name (as shown)', 'Product URL / screen', 'Product ID (if in URL)',
            'Brand', 'Manufacturer', 'Composition', 'Strength', 'Dosage form', 'Pack size', 'Rx required (shown)', 'Category',
            'MRP ₹', 'Selling price ₹', 'Discount %', 'Availability', 'Add-to-cart enabled', 'ETA text (exact)',
            'Max qty shown', 'Substitutes shown', 'Screenshot', 'Notes'],
           widths={'Product name (as shown)': 34, 'Composition': 28})
for r in range(5, 2001):
    ws[f'Q{r}'] = f'=IF(AND(ISNUMBER(O{r}),ISNUMBER(P{r}),O{r}>0),ROUND((O{r}-P{r})/O{r}*100,1),"")'
dropdown(ws, 'R', ['in_stock', 'out_of_stock', 'unknown'])
dropdown(ws, 'S', ['yes', 'no', 'unknown'])
dropdown(ws, 'M', ['yes', 'no', 'not_shown'])

# 4. SKU Availability
cats = ['Prescription medicines', 'OTC medicines', 'Vitamins & supplements', 'Personal care', 'Baby care',
        'Diabetes care', 'Medical devices', 'Ayurveda', 'Other']
ws = sheet('SKU Availability', 'Calculated from Master SKU Catalogue (column N must use these category names).',
           ['Category', 'SKUs observed', 'In stock', 'Out of stock', 'Unknown', 'Observed availability %', 'Note'],
           [(c,) for c in cats])
for i, _ in enumerate(cats):
    r = 5 + i
    rng = "'Master SKU Catalogue'!$N$5:$N$5000"
    av = "'Master SKU Catalogue'!$R$5:$R$5000"
    ws[f'B{r}'] = f'=COUNTIF({rng},A{r})'
    ws[f'C{r}'] = f'=COUNTIFS({rng},A{r},{av},"in_stock")'
    ws[f'D{r}'] = f'=COUNTIFS({rng},A{r},{av},"out_of_stock")'
    ws[f'E{r}'] = f'=COUNTIFS({rng},A{r},{av},"unknown")'
    ws[f'F{r}'] = f'=IFERROR(ROUND(C{r}/(C{r}+D{r})*100,1),"n/a")'
    ws[f'G{r}'] = 'Denominator excludes unknown status'

# 5. Medicine Details
sheet('Medicine Details', 'Only what the product page shows. Do not infer equivalence between brands.',
      ['Obs ID (from catalogue)', 'Product name', 'Salt composition', 'Therapeutic class (shown)', 'Storage', 'Warnings',
       'Expiry shown', 'Batch shown', 'Substitutes shown', 'Alt strengths', 'Alt pack sizes', 'Screenshot'])

# 6. Pricing and Discounts
ws = sheet('Pricing and Discounts', 'Calculated from Master SKU Catalogue.',
           ['Category', 'SKUs with MRP and price', 'Average discount %', 'Max discount %', 'Selling > MRP (flag)'],
           [(c,) for c in cats])
for i, _ in enumerate(cats):
    r = 5 + i
    n, q = "'Master SKU Catalogue'!$N$5:$N$5000", "'Master SKU Catalogue'!$Q$5:$Q$5000"
    ws[f'B{r}'] = f'=COUNTIFS({n},A{r},{q},">=-1000")'
    ws[f'C{r}'] = f'=IFERROR(ROUND(AVERAGEIFS({q},{n},A{r}),1),"n/a")'
    ws[f'D{r}'] = f'=IFERROR(_xlfn.MAXIFS({q},{n},A{r}),"n/a")'
    ws[f'E{r}'] = f'=COUNTIFS({n},A{r},{q},"<0")'

# 7. Delivery ETA
ws = sheet('Delivery ETA', 'Exact text as displayed. Slots: morning 8–10, afternoon 12–2, evening 6–8, night 9–11 (IST). A promise is not a measured delivery time.',
           ['Obs ID', 'Date', 'Time (IST)', 'Slot', 'ETA text (exact)', 'Where shown', 'Address PIN', 'Express option', 'Scheduled option', 'Screenshot', 'Notes'],
           [('ETA-0', OBS, '', '', '15 minutes (claimed)', 'structured data / FAQ', '', '', '', '', 'Marketing claim, not a displayed checkout ETA')])
dropdown(ws, 'D', ['morning', 'afternoon', 'evening', 'night'])
dropdown(ws, 'F', ['home', 'category', 'product', 'cart', 'checkout'])

# 8. Checkout Fee Matrix
fees = ['Delivery fee', 'Handling fee', 'Platform fee', 'Packaging fee', 'Small-cart fee', 'Surge / late-night fee',
        'Priority fee', 'Rx processing fee', 'COD fee', 'Payment fee', 'Free-delivery threshold', 'Taxes shown']
sheet('Checkout Fee Matrix', 'Fill from cart/checkout screens (stop before payment).',
      ['Fee', 'Seen?', 'Amount / rule (exact text)', 'Condition (basket value, time)', 'Basket ID', 'Screenshot', 'Notes'],
      [(f,) for f in fees])

# 9. Basket Experiments
targets = [99, 199, 299, 399, 499, 599, 699, 799, 999, 1499, 1999, 2999]
ws = sheet('Basket Experiments',
           'Build each basket in the cart only; never pay or place the order. Total check = 0 means the arithmetic matches.',
           ['Basket ID', 'Target ₹', 'Items (name × qty)', 'Basket MRP ₹', 'Selling subtotal ₹', 'Product discount ₹',
            'Coupon ₹', 'Delivery ₹', 'Handling ₹', 'Platform ₹', 'Other ₹', 'Total payable (shown) ₹', 'Computed total ₹',
            'Total check ₹', 'Fees % of subtotal', 'ETA text', 'Free delivery?', 'Mix', 'Observed at', 'Screenshot'],
           [(f'B{i + 1:02d}', t) for i, t in enumerate(targets)], {'Items (name × qty)': 40})
for r in range(5, 101):
    ws[f'M{r}'] = f'=IF(ISNUMBER(E{r}),E{r}-N(G{r})+N(H{r})+N(I{r})+N(J{r})+N(K{r}),"")'
    ws[f'N{r}'] = f'=IF(AND(ISNUMBER(L{r}),ISNUMBER(M{r})),ROUND(L{r}-M{r},2),"")'
    ws[f'O{r}'] = f'=IF(AND(ISNUMBER(E{r}),E{r}>0),ROUND((N(H{r})+N(I{r})+N(J{r})+N(K{r}))/E{r}*100,1),"")'
dropdown(ws, 'R', ['medicines', 'otc', 'wellness', 'mixed'], last=100)

# 10. Promotions
sheet('Promotions', 'Offers visible to an ordinary customer.',
      ['Promo', 'Type', 'Code (if public)', 'Discount rule', 'Min basket ₹', 'Cap ₹', 'Category', 'Validity', 'Where seen', 'Observed at', 'Screenshot'],
      [('Discounts, cashback, coupons', 'claimed', '', '"Download the app to get the best deals" (FAQ)', '', '', '', '', 'FAQ', OBS, '')])

# 11. Search Tests
queries = [('brand', 'Dolo 650'), ('brand', 'Pan 40'), ('brand', 'Crocin Advance'), ('generic', 'paracetamol'),
           ('salt', 'pantoprazole 40'), ('salt', 'metformin 500'), ('typo', 'paracetmol'), ('typo', 'azithromicin'),
           ('no-space', 'dolo650'), ('partial', 'thyro'), ('symptom', 'fever'), ('symptom', 'acidity'),
           ('brand+strength', 'Telma 40'), ('brand+pack', 'Ecosprin 75 strip of 14'), ('device', 'Accu-Chek Active strips'),
           ('wellness', 'Ensure'), ('personal care', 'Whisper'), ('multi', 'dolo 650, pan 40'), ('cold chain', 'insulin'),
           ('Hindi', 'bukhar ki dawai')]
ws = sheet('Search Tests', 'Type the query in the site/app search at the fixed address. Results are not medical advice.',
           ['Query type', 'Query', 'Results (count)', 'Exact match position', 'Top 5 results', 'Autocomplete / suggestions',
            'Typo handled?', 'Out-of-stock shown?', 'Alternatives offered', 'Filters / sort', 'Observed at', 'Screenshot'],
           queries, {'Top 5 results': 45})
dropdown(ws, 'G', ['yes', 'no', 'n/a'], last=200)

# 12. Prescription Workflows
rx = ['Prescription upload', 'Multiple prescriptions', 'Multi-medicine order', 'Pharmacist review messaging',
      'Rx-to-cart conversion', 'Substitution consent', 'Unavailable medicine handling', 'Partial fulfilment',
      'Split deliveries', 'Refill reminders', 'Repeat orders', 'Family profiles', 'Saved prescriptions', 'WhatsApp ordering']
ws = sheet('Prescription Workflows', 'Never upload a real prescription or patient data. Status: observed / claimed / not_found / untestable.',
           ['Feature', 'Status', 'What is shown (exact)', 'Source', 'Observed at', 'Screenshot'],
           [(f,) for f in rx])
ws['B5'], ws['C5'], ws['D5'], ws['E5'] = 'claimed', 'Upload prescription on the app (FAQ, HowTo)', SRC, OBS
ws['B17'], ws['C17'], ws['D17'], ws['E17'] = 'claimed', '/saved-prescriptions page listed in sitemap', 'https://plazza.in/sitemap.xml', OBS
dropdown(ws, 'B', STATUS, last=100)

# 13. Membership and Loyalty
mem = ['Membership plan', 'Price', 'Duration', 'Free delivery', 'Fee waivers', 'Extra discount', 'Family cover',
       'Renewal / cancellation terms', 'Referral credit', 'Loyalty points', 'Wallet / cashback']
ws = sheet('Membership and Loyalty', 'Pinzzo benchmark column is our plan, not a competitor fact.',
           ['Feature', 'Status', 'Plazza (exact)', 'Pinzzo benchmark', 'Source', 'Observed at'],
           [(f, '', '', '₹149 / 90 days' if f == 'Price' else '') for f in mem])
ws['B15'], ws['C15'], ws['E15'], ws['F15'] = 'claimed', '"cashback offers" (FAQ)', SRC, OBS
dropdown(ws, 'B', STATUS, last=100)

# 14. WhatsApp and Ordering Channels
sheet('WhatsApp and Ordering Channels', 'Public entry points only; no automated or bulk messaging.',
      ['Channel', 'Status', 'Entry point', 'What it offers', 'Hours', 'Source', 'Observed at'],
      [('Website', 'observed', 'https://plazza.in', 'Catalogue, cart (client-side)', '24/7 claimed', SRC, OBS),
       ('Android app', 'observed', 'https://play.google.com/store/apps/details?id=com.plazza.app', '', '', SRC, OBS),
       ('iOS app', 'observed', 'https://apps.apple.com/in/app/plazza-medicines-in-minutes/id6738200177', '', '', SRC, OBS),
       ('Email', 'claimed', 'support@plazza.in; founders@plazza.in', '', '', SRC, OBS),
       ('Phone', 'not_found', 'Structured data shows a placeholder "+91-XXXXX-XXXXX"', '', '', SRC, OBS),
       ('WhatsApp', 'pending', '', '', '', '', ''),
       ('Instagram / LinkedIn', 'observed', 'instagram.com/plazza.in; linkedin.com/company/plazza-club', '', '', SRC, OBS)],
      {'Entry point': 55})

# 15. Serviceability
sheet('Serviceability', 'Use ONE fixed Bengaluru address for the whole study. Do not bypass location checks.',
      ['Field', 'Value', 'Kind', 'Source', 'Observed at'],
      [('Fixed study address', '[choose a Bengaluru address, e.g. near Bellandur 560103]', 'input', '', ''),
       ('Serviceable at study address?', '', '', '', ''),
       ('Location selection mechanism', '', '', '', ''),
       ('Claimed service area', 'Bengaluru', 'claimed', SRC, OBS),
       ('Plazza office coordinates', '12.9327764, 77.6834966 (Bellandur)', 'claimed', SRC, OBS),
       ('Operating hours', '24/7', 'claimed', SRC, OBS),
       ('Overlap with Pinzzo (Gurugram)', 'None observed: Plazza claims Bengaluru only', 'observed', SRC, OBS)],
      {'Value': 55})

# 16. Customer Experience
steps = ['Landing', 'Location selection', 'Login (OTP)', 'Search / browse', 'Product page', 'Add to cart', 'Cart',
         'Address', 'Prescription step', 'Coupon', 'Checkout summary (stop here)']
sheet('Customer Experience', 'Journey to the last screen before payment. Never place an order.',
      ['Step', 'Required?', 'What happens (exact)', 'Friction / blockers', 'Time taken (s)', 'Screenshot'],
      [(s,) for s in steps])

# 17. Policies and Compliance
pol = [('Legal entity', 'PLAZZA HEALTH PRIVATE LIMITED', 'claimed', SRC),
       ('Address', 'Bellandur, Bengaluru 560103', 'claimed', SRC),
       ('Drug licence number', '', 'pending', 'not found in page HTML'),
       ('Terms', 'Text loads client-side; capture manually', 'pending', 'https://plazza.in/terms'),
       ('Privacy policy', 'Text loads client-side; capture manually', 'pending', 'https://plazza.in/privacy-policy'),
       ('Cancellation & refund', 'Text loads client-side; capture manually', 'pending', 'https://plazza.in/cancellation-refund'),
       ('Grievance officer', '', 'pending', ''),
       ('Medicines genuine / licensed claim', '"licensed pharmacy", "100% genuine" (FAQ)', 'claimed', SRC)]
sheet('Policies and Compliance', 'Record source URL and date for each item.',
      ['Item', 'Value / summary', 'Status', 'Source', 'Observed at'],
      [(a, b, c, d, OBS) for a, b, c, d in pol], {'Value / summary': 55, 'Source': 40})

# 18. Change Detection
sheet('Change Detection', 'One row per repeat observation of the same product. A single snapshot is not a trend.',
      ['Snapshot ID', 'Observed at', 'Address PIN', 'Product (Obs ID)', 'MRP ₹', 'Selling ₹', 'Availability', 'ETA text',
       'Fees text', 'Promo', 'Changed vs previous?', 'Notes'])

# 19. Pinzzo Comparison
cmp = [('Delivery ETA', '15 min (claimed)', '30–40 min'),
       ('Operating hours', '24/7 (claimed)', 'TBD'),
       ('Assortment', '20,000–100,000+ (claims conflict)', 'Medicine-first, localised'),
       ('SKU availability', 'pending manual sample', 'Localisation-led'),
       ('Minimum order', 'None (claimed)', 'TBD'),
       ('Delivery fees', '₹0 in Offer data; checkout fees pending', '₹25 under ₹499; ₹5 handling; ₹15 night'),
       ('Discounts', 'Coupons/cashback claimed; depth pending', '10% model assumption'),
       ('Membership', 'pending', '₹149 / 90 days'),
       ('Prescription ordering', 'Upload in app (claimed); saved prescriptions page exists', 'Multi-prescription'),
       ('WhatsApp ordering', 'pending', 'Planned'),
       ('Split shipments', 'pending', 'Avoid'),
       ('Geography', 'Bengaluru', 'Gurugram (no overlap today)'),
       ('Customer retention', 'Public evidence only: app 4.1 from 31 ratings (their page)', 'Refill-led')]
sheet('Pinzzo Comparison', 'Competitor column = observed or claimed facts only. Opportunity is our analysis.',
      ['Dimension', 'Plazza observed / claimed', 'Pinzzo benchmark', 'Opportunity'], cmp,
      {'Plazza observed / claimed': 50, 'Pinzzo benchmark': 35, 'Opportunity': 50})

# 20. Data Quality and Limitations
lim = [('Automated product, price and stock collection', 'Not performed', 'Product data loads only from /api/, which plazza.in/robots.txt disallows for all crawlers. Collected manually instead.'),
       ('Policy pages', 'Manual only', 'Terms, privacy and refund text is rendered client-side from the same API.'),
       ('Assortment claims', 'Conflicting', 'Page structured data says 20,000+, 35,000+ and 100,000+ in different blocks.'),
       ('Structured data freshness', 'Stale fields', 'priceValidUntil 2025-12-31; phone is a placeholder.'),
       ('App rating', 'Unverified', '4.1 / 31 ratings is Plazza\'s own markup, not checked on the stores.'),
       ('Geography', 'Different market', 'Plazza claims Bengaluru only; findings benchmark the model, not Gurugram prices.'),
       ('Inventory quantities', 'Out of scope', 'Not inferred from add-to-cart or any probing.'),
       ('Orders, payments, real prescriptions', 'Out of scope', 'No orders placed, no payments, no medical documents uploaded.')]
sheet('Data Quality and Limitations', 'Register of every constraint and exclusion.',
      ['Area', 'Status', 'Detail'], lim, {'Area': 40, 'Status': 18, 'Detail': 90})

wb.save(OUT)
print('saved', OUT)
